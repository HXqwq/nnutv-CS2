# -*- coding: utf-8 -*-
"""NNUTV CS2 比赛数据填报表 -> data/matches.json 自动导入脚本。

用法（在本仓库根目录）：
  python scripts/import-sheet.py <填好的表格.xlsx>               # 自动分配 m{最大+1}
  python scripts/import-sheet.py <填好的表格.xlsx> --match-id m35 # 指定比赛 id（重跑覆盖）
  python scripts/import-sheet.py <填好的表格.xlsx> --set G7=95    # 修正基础数据单元格笔误（可多次）
  python scripts/import-sheet.py <填好的表格.xlsx> --force        # 校验不自洽时仍强制导入（如自杀导致的 h2h≠死亡）

行为：
  1. 读「比赛信息 / 基础数据 / 对位数据」三张表
  2. 重跑全部自洽校验（K/D/A 半场加和、h2h 行列累加）
  3. 昵称匹配 players.json（精确 -> 去空格小写），未命中则新增 p{max+1}
  4. 队名匹配 teams.json，未命中则新增 t{max+1}（shortName=队名，country 留空）
  5. 按 id 覆盖合并写入 matches.json（JSON 格式与 node 脚本输出一致）
"""
import argparse
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

METRICS = ["kills", "deaths", "assists", "adr", "hs", "rating"]
N_PLAYERS = 10


def as_num(v, force_int=False):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if not v:
            return None
        v = float(v)
    if force_int or (isinstance(v, float) and v.is_integer()):
        return int(v)
    return v


def norm(s):
    return re.sub(r"\s+", "", str(s)).lower()


def load_json(name):
    return json.loads((DATA / name).read_text(encoding="utf-8"))


def save_json(name, obj):
    text = json.dumps(obj, ensure_ascii=False, indent=2) + "\n"
    (DATA / name).write_text(text, encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--match-id")
    ap.add_argument("--set", action="append", default=[], metavar="CELL=VALUE",
                    help="覆盖基础数据单元格，如 --set G7=95")
    ap.add_argument("--force", action="store_true", help="校验失败仍导入")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    info, base, h2h_ws = wb["比赛信息"], wb["基础数据"], wb["对位数据"]

    # ---- 人工覆盖（修正表格笔误） ----
    for kv in args.set:
        ref, val = kv.split("=", 1)
        base[ref.strip().upper()] = float(val) if "." in val else int(val)
        print(f"[override] 基础数据!{ref.strip().upper()} = {val}")

    # ---- 比赛信息 ----
    fields = {}
    for r in range(4, 17):
        label = info.cell(row=r, column=1).value
        if label:
            fields[str(label).strip()] = info.cell(row=r, column=2).value
    date = fields["日期"]
    date = date.strftime("%Y-%m-%d") if isinstance(date, (datetime.datetime, datetime.date)) else str(date)
    match_info = {
        "date": date,
        "event": str(fields["赛事名称"]).strip(),
        "stage": str(fields["赛段"]).strip(),
        "format": str(fields["赛制"]).strip(),
        "teamA": str(fields["A队队名"]).strip(),
        "teamB": str(fields["B队队名"]).strip(),
        "scoreA": as_num(fields["A队得分"], force_int=True),
        "scoreB": as_num(fields["B队得分"], force_int=True),
    }

    # ---- 基础数据 ----
    players = []
    for r in range(4, 4 + N_PLAYERS):
        nick = base.cell(row=r, column=3).value
        if not nick:
            continue
        row = {"row": r, "nick": str(nick), "side": str(base.cell(row=r, column=2).value).strip().upper()}
        for block, col0 in (("", 4), ("ct_", 10), ("t_", 16)):
            vals = {k: as_num(base.cell(row=r, column=col0 + j).value, force_int=(k != "rating"))
                    for j, k in enumerate(METRICS)}
            if block:
                row[block + "stats"] = vals
            else:
                row.update(vals)
        players.append(row)
    if len(players) != N_PLAYERS:
        sys.exit(f"FAIL 基础数据只填了 {len(players)} 人，应为 {N_PLAYERS}")
    a_rows = [p for p in players if p["side"] == "A"]
    b_rows = [p for p in players if p["side"] == "B"]
    if len(a_rows) != 5 or len(b_rows) != 5:
        sys.exit(f"FAIL 队伍侧分布 A={len(a_rows)} B={len(b_rows)}，应各 5 人")

    # ---- 对位数据：定位 ①② 两个矩阵块 ----
    blocks = {}
    for r in range(1, h2h_ws.max_row + 1):
        title = h2h_ws.cell(row=r, column=1).value
        if isinstance(title, str) and title.strip()[:1] in ("①", "②"):
            blocks[title.strip()[:1]] = r
        if len(blocks) == 2:
            break
    if len(blocks) != 2:
        sys.exit("FAIL 对位数据表找不到 ①/② 两个矩阵块")

    def read_matrix(top):
        first = top + 2
        grid = {}
        for i, ap_ in enumerate(a_rows):
            for j, bp in enumerate(b_rows):
                grid[(ap_["nick"], bp["nick"])] = as_num(
                    h2h_ws.cell(row=first + i, column=2 + j).value, force_int=True)
        return grid

    kill_ab = read_matrix(blocks["①"])   # 行(A) 击杀 列(B)
    death_ab = read_matrix(blocks["②"])  # 行(A) 被列(B) 击杀

    # ---- 校验 ----
    problems = []
    for p in players:
        for k in ("kills", "deaths", "assists"):
            if p[k] != p["ct_stats"][k] + p["t_stats"][k]:
                problems.append(f"半场加和: {p['nick']} {k} 全场{p[k]} != CT{p['ct_stats'][k]}+T{p['t_stats'][k]}")
    for p in a_rows:
        s1 = sum(kill_ab[(p["nick"], bp["nick"])] for bp in b_rows)
        if s1 != p["kills"]:
            problems.append(f"对位① 行合计: {p['nick']} {s1} != 全场击杀 {p['kills']}")
        s2 = sum(death_ab[(p["nick"], bp["nick"])] for bp in b_rows)
        if s2 != p["deaths"]:
            problems.append(f"对位② 行合计: {p['nick']} {s2} != 全场死亡 {p['deaths']}（自杀/坠亡等非对位死亡可 --force）")
    for bp in b_rows:
        c1 = sum(kill_ab[(ap_["nick"], bp["nick"])] for ap_ in a_rows)
        if c1 != bp["deaths"]:
            problems.append(f"对位① 列合计: {bp['nick']} {c1} != 全场死亡 {bp['deaths']}")
        c2 = sum(death_ab[(ap_["nick"], bp["nick"])] for ap_ in a_rows)
        if c2 != bp["kills"]:
            problems.append(f"对位② 列合计: {bp['nick']} {c2} != 全场击杀 {bp['kills']}")
    if problems:
        for x in problems:
            print("WARN", x)
        if not args.force:
            sys.exit("FAIL 校验未通过（如属自杀等合理情况请加 --force）")

    # ---- 队伍解析 ----
    teams = load_json("teams.json")
    def team_ref(name):
        hit = next((t for t in teams if t["name"] == name), None)
        if hit:
            return hit["name"]
        hit = next((t for t in teams if norm(t.get("shortName", "")) == norm(name)), None)
        if hit:
            return hit["name"]
        nid = "t" + str(max((int(t["id"][1:]) for t in teams), default=0) + 1)
        teams.append({"id": nid, "name": name, "shortName": name, "country": ""})
        print(f"[teams] 新增 {nid} {name}")
        return name

    team_a, team_b = team_ref(match_info["teamA"]), team_ref(match_info["teamB"])
    match_info["teamA"], match_info["teamB"] = team_a, team_b

    # ---- 选手解析 ----
    plist = load_json("players.json")
    exact = {p["nickname"]: p for p in plist}
    loose = {norm(p["nickname"]): p for p in plist}
    created = []

    def player_ref(nick, team):
        hit = exact.get(nick) or loose.get(norm(nick))
        if hit:
            return hit["id"]
        nid = "p" + str(max((int(p["id"][1:]) for p in plist), default=0) + 1)
        plist.append({"id": nid, "nickname": nick, "realName": "", "team": team,
                      "role": "", "joined": date[:7]})
        created.append(nid)
        print(f"[players] 新增 {nid} {nick} (team={team})")
        return nid

    for p in players:
        p["id"] = player_ref(p["nick"], team_a if p["side"] == "A" else team_b)

    # ---- 组装 match ----
    ms = load_json("matches.json")
    mid = args.match_id
    if not mid:
        mid = "m" + str(max((int(m["id"][1:]) for m in ms if re.fullmatch(r"m\d+", m["id"])), default=0) + 1)

    stats = []
    for p in players:
        opp = b_rows if p["side"] == "A" else a_rows
        h2h = {}
        for o in opp:
            if p["side"] == "A":
                h2h[o["id"]] = {"killed": kill_ab[(p["nick"], o["nick"])],
                                "death": death_ab[(p["nick"], o["nick"])]}
            else:
                h2h[o["id"]] = {"killed": death_ab[(o["nick"], p["nick"])],
                                "death": kill_ab[(o["nick"], p["nick"])]}
        stats.append({
            "playerId": p["id"], "side": p["side"],
            "kills": p["kills"], "deaths": p["deaths"], "assists": p["assists"],
            "adr": p["adr"], "hs": p["hs"], "rating": p["rating"],
            "ctStats": p["ct_stats"], "tStats": p["t_stats"],
            "multiKills": None, "mvp": 0, "rws": None, "we": None,
            "performance": None, "h2h": h2h,
        })

    match = {**match_info, "id": mid, "maps": [{
        "map": str(fields["地图名"]).strip(),
        "scoreA": match_info["scoreA"], "scoreB": match_info["scoreB"],
        "stats": stats,
    }]}

    idx = next((i for i, m in enumerate(ms) if m["id"] == mid), None)
    if idx is None:
        ms.append(match)
        action = "新增"
    else:
        ms[idx] = match
        action = "覆盖"

    # ---- 写盘 ----
    save_json("players.json", plist)
    save_json("teams.json", teams)
    save_json("matches.json", ms)
    print(f"OK {action} {mid}: {team_a} {match_info['scoreA']}:{match_info['scoreB']} {team_b} "
          f"| {match_info['date']} {match_info['event']} {match_info['stage']} {match_info['format']} "
          f"| 地图 {match['maps'][0]['map']} | 选手行 {len(stats)}")
    if created:
        print("新选手:", ", ".join(f"{pid}={next(p['nick'] for p in players if p['id'] == pid)}"
                                   for pid in created))


if __name__ == "__main__":
    main()
